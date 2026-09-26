from __future__ import annotations

import json
import shutil
import zipfile
from pathlib import Path

import pytest
import yaml
from openpyxl import load_workbook

from cyber_assess.artifacts import create_final_artifacts
from cyber_assess.errors import InvalidInputError
from cyber_assess.reporting import create_deliverable_inputs
from cyber_assess.snapshots import IntegrityVerificationError, load_snapshot

SNAPSHOT_ID = "SNAP-20260926T012311-1F44808C"


def _workspace() -> Path:
    return Path(__file__).resolve().parents[2]


def _copy_frozen_workspace(destination: Path) -> Path:
    source_workspace = _workspace()
    snapshot = load_snapshot(source_workspace, SNAPSHOT_ID)
    for source in snapshot.source_hashes:
        source_path = source_workspace / source.path
        target_path = destination / source.path
        target_path.parent.mkdir(parents=True, exist_ok=True)
        shutil.copyfile(source_path, target_path)
    snapshot_path = source_workspace / "assessment" / "snapshots" / f"{SNAPSHOT_ID}.yaml"
    target_snapshot = destination / "assessment" / "snapshots" / snapshot_path.name
    target_snapshot.parent.mkdir(parents=True, exist_ok=True)
    shutil.copyfile(snapshot_path, target_snapshot)
    return destination


def _bundle_bytes(bundle: Path) -> dict[str, bytes]:
    return {path.name: path.read_bytes() for path in sorted(bundle.iterdir())}


def test_deliverable_inputs_are_deterministic_reconciled_and_allowlisted(
    tmp_path: Path,
) -> None:
    first = create_deliverable_inputs(_workspace(), SNAPSHOT_ID, output_root=tmp_path / "first")
    second = create_deliverable_inputs(_workspace(), SNAPSHOT_ID, output_root=tmp_path / "second")

    assert _bundle_bytes(first) == _bundle_bytes(second)
    assert set(_bundle_bytes(first)) == {
        "claims-evidence-index.json",
        "executive-report-input.md",
        "presentation-input.md",
        "reconciliation.json",
        "remediation-roadmap-input.json",
        "reporting-manifest.json",
        "workbook-input.json",
    }

    reconciliation = json.loads((first / "reconciliation.json").read_text(encoding="utf-8"))
    assert reconciliation["record_counts"] == {
        "actions": 1,
        "controls": 20,
        "evidence": 4,
        "findings": 1,
        "retests": 1,
        "workpapers": 1,
    }
    assert reconciliation["control_coverage"] == {
        "assessed": 1,
        "state": "partial",
        "total": 20,
        "unassessed": 19,
        "workpapers_missing": 19,
        "workpapers_present": 1,
    }
    assert reconciliation["snapshot_alignment"]["unmatched_references"] == 0

    index = json.loads((first / "claims-evidence-index.json").read_text(encoding="utf-8"))
    assert index["claims"][0]["workpaper_id"] == "WP-IAM-02-S01"
    assert all(item["classification"] == "sanitized" for item in index["evidence"])
    serialized = b"\n".join(_bundle_bytes(first).values()).lower()
    assert b"password" not in serialized
    assert b"backup.sql" not in serialized
    assert b"fixtures/evidence" not in serialized
    assert b"query_or_action" not in serialized


def test_tampered_snapshot_source_is_not_published(tmp_path: Path) -> None:
    workspace = _copy_frozen_workspace(tmp_path / "workspace")
    (workspace / "assessment" / "engagement.yaml").write_text(
        "schema_version: '1.0'\nid: ENG-TAMPERED\n", encoding="utf-8"
    )
    output_root = tmp_path / "outputs"

    with pytest.raises(IntegrityVerificationError, match="hash mismatch"):
        create_deliverable_inputs(workspace, SNAPSHOT_ID, output_root=output_root)

    assert not (output_root / SNAPSHOT_ID).exists()


def test_snapshot_id_reconciliation_failure_is_not_published(tmp_path: Path) -> None:
    workspace = _copy_frozen_workspace(tmp_path / "workspace")
    snapshot_path = workspace / "assessment" / "snapshots" / f"{SNAPSHOT_ID}.yaml"
    snapshot_data = yaml.safe_load(snapshot_path.read_text(encoding="utf-8"))
    snapshot_data["baseline_ids"] = ["FIND-NOT-FROZEN"]
    snapshot_path.write_text(
        yaml.safe_dump(snapshot_data, sort_keys=False), encoding="utf-8", newline="\n"
    )
    output_root = tmp_path / "outputs"

    with pytest.raises(InvalidInputError, match="baseline IDs"):
        create_deliverable_inputs(workspace, SNAPSHOT_ID, output_root=output_root)

    assert not (output_root / SNAPSHOT_ID).exists()


def test_existing_bundle_is_not_overwritten(tmp_path: Path) -> None:
    output_root = tmp_path / "outputs"
    created = create_deliverable_inputs(_workspace(), SNAPSHOT_ID, output_root=output_root)
    original = _bundle_bytes(created)

    with pytest.raises(FileExistsError, match="already exists"):
        create_deliverable_inputs(_workspace(), SNAPSHOT_ID, output_root=output_root)

    assert _bundle_bytes(created) == original


def test_final_artifacts_are_deterministic_formatted_and_honest(tmp_path: Path) -> None:
    first = create_final_artifacts(_workspace(), SNAPSHOT_ID, output_root=tmp_path / "first")
    second = create_final_artifacts(_workspace(), SNAPSHOT_ID, output_root=tmp_path / "second")

    assert _bundle_bytes(first) == _bundle_bytes(second)
    assert set(_bundle_bytes(first)) == {
        "artifact-manifest.json",
        "assessment-workbook.xlsx",
        "engagement-brief.html",
        "engagement-brief.md",
        "executive-report.html",
        "executive-report.md",
        "pdf-status.json",
        "presentation.html",
        "reconciliation.json",
        "remediation-roadmap.csv",
        "remediation-roadmap.json",
    }

    workbook = load_workbook(first / "assessment-workbook.xlsx", read_only=False)
    assert workbook.sheetnames == [
        "Cover",
        "Scope",
        "Assets",
        "Risk Register",
        "Control Register",
        "Workpapers",
        "Evidence Index",
        "Findings",
        "Remediation",
        "Retests",
        "NIST Mapping",
        "RMiT Mapping",
        "Reconciliation",
        "Limitations",
    ]
    for sheet in workbook.worksheets[1:]:
        assert sheet.freeze_panes == "A2"
        assert sheet.auto_filter.ref
        assert sheet["A1"].font.bold
    evidence_values = {
        str(cell.value)
        for row in workbook["Evidence Index"].iter_rows()
        for cell in row
        if cell.value
    }
    assert "EVID-S01-BASELINE" in evidence_values
    assert "EVID-S01-RETEST" in evidence_values

    presentation = (first / "presentation.html").read_text(encoding="utf-8")
    assert presentation.count('class="slide"') == 7
    assert presentation.count('class="speaker-notes"') == 7
    assert "PDF status:</strong> Pending" in presentation
    pdf_status = json.loads((first / "pdf-status.json").read_text(encoding="utf-8"))
    assert pdf_status == {
        "claimed_pdf_artifacts": [],
        "reason": "No locked PDF renderer is available; print-ready HTML is provided.",
        "status": "pending",
    }

    text_payload = b"\n".join(
        path.read_bytes() for path in first.iterdir() if path.suffix != ".xlsx"
    ).lower()
    with zipfile.ZipFile(first / "assessment-workbook.xlsx") as archive:
        workbook_payload = b"\n".join(
            archive.read(name) for name in archive.namelist() if name.endswith(".xml")
        ).lower()
    combined = text_payload + workbook_payload
    assert b"fixtures/evidence" not in combined
    assert b"query_or_action" not in combined
    assert b"backup.sql" not in combined


def test_final_artifacts_reject_tampered_sources_without_publication(tmp_path: Path) -> None:
    workspace = _copy_frozen_workspace(tmp_path / "workspace")
    (workspace / "assessment" / "engagement.yaml").write_text(
        "schema_version: '1.0'\nid: ENG-TAMPERED\n", encoding="utf-8"
    )
    output_root = tmp_path / "outputs"

    with pytest.raises(IntegrityVerificationError, match="hash mismatch"):
        create_final_artifacts(workspace, SNAPSHOT_ID, output_root=output_root)

    assert not (output_root / SNAPSHOT_ID).exists()

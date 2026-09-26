"""Deterministic safe artifact rendering from a verified assessment snapshot."""

from __future__ import annotations

import csv
import hashlib
import html
import io
import os
import shutil
import tempfile
import zipfile
from collections.abc import Iterable, Sequence
from pathlib import Path

from markdown import markdown  # type: ignore[import-untyped]
from openpyxl import Workbook  # type: ignore[import-untyped]
from openpyxl.styles import Alignment, Font, PatternFill  # type: ignore[import-untyped]
from openpyxl.worksheet.worksheet import Worksheet  # type: ignore[import-untyped]

from cyber_assess.models import ExportSnapshot
from cyber_assess.reporting import (
    _FrozenRecords,
    _iso,
    _json_text,
    _limitations,
    _load_frozen_records,
    _reconcile,
)
from cyber_assess.snapshots import load_snapshot, verify_snapshot

_BLUE = "17365D"
_ZIP_TIME = (1980, 1, 1, 0, 0, 0)
_ARTIFACT_NAMES = (
    "assessment-workbook.xlsx",
    "engagement-brief.html",
    "engagement-brief.md",
    "executive-report.html",
    "executive-report.md",
    "pdf-status.json",
    "presentation.html",
    "reconciliation.json",
    "remediation-roadmap.csv",
    "remediation-roadmap.json",
)


def _final_limitations(records: _FrozenRecords, snapshot: ExportSnapshot) -> list[str]:
    values = _limitations(records, snapshot)
    values[4] = (
        "PDF artifacts, when present, are local derivatives with a separate render manifest "
        "and require visual review before distribution."
    )
    return values


def _coverage(reconciliation: dict[str, object]) -> dict[str, object]:
    value = reconciliation["control_coverage"]
    if not isinstance(value, dict):
        raise TypeError("control coverage must be a mapping")
    return value


def _md(value: object) -> str:
    return str(value).replace("\\", "\\\\").replace("|", "\\|").replace("\n", " ")


def _html_document(title: str, body_markdown: str) -> str:
    body = markdown(body_markdown, extensions=["tables"])
    return f"""<!doctype html>
<html lang="en"><head><meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>{html.escape(title)}</title><style>
@page {{ size: A4; margin: 18mm; }}
body {{ color:#172033; font:15px/1.55 Arial,sans-serif; margin:0 auto;
max-width:980px; padding:32px }} h1,h2 {{ color:#17365d }}
table {{ border-collapse:collapse; width:100% }} th,td {{ border:1px solid #aeb9c5;
padding:7px; text-align:left }} th {{ background:#eef4f8 }}
.pdf-status {{ border:2px solid #9a6700; background:#fff8c5; padding:12px }}
@media print {{ body {{ max-width:none; padding:0 }} .pdf-status {{ display:none }} }}
</style></head><body><p class="pdf-status"><strong>PDF status:</strong> Pending.
This HTML is print-ready, but no PDF is claimed or included.</p>{body}</body></html>
"""


def _brief_markdown(
    snapshot: ExportSnapshot,
    records: _FrozenRecords,
    reconciliation: dict[str, object],
    limitations: Sequence[str],
) -> str:
    coverage = _coverage(reconciliation)
    lines = [
        f"# {_md(records.engagement.name)}",
        "",
        "## Engagement brief",
        "",
        f"- Snapshot: `{snapshot.id}`",
        f"- Frozen at: `{_iso(snapshot.created_at)}`",
        f"- Assessment period: {records.engagement.assessment_period.start} "
        f"to {records.engagement.assessment_period.end}",
        f"- Review mode: `{records.engagement.review_mode}`",
        f"- Workpaper coverage: {coverage['workpapers_present']} of {coverage['total']}",
        f"- Executed assessment conclusions: {coverage['assessed']} of {coverage['total']} "
        f"(`{coverage['state']}`)",
        "",
        "## Business context",
        "",
        _md(records.engagement.business_context),
        "",
        "## Scope",
        "",
        *[f"- {_md(item)}" for item in records.engagement.scope],
        "",
        "## Exclusions",
        "",
        *[f"- {_md(item)}" for item in records.engagement.exclusions],
        "",
        "## Limitations",
        "",
        *[f"- {_md(item)}" for item in limitations],
    ]
    return "\n".join(lines) + "\n"


def _executive_markdown(
    snapshot: ExportSnapshot,
    records: _FrozenRecords,
    reconciliation: dict[str, object],
    limitations: Sequence[str],
) -> str:
    coverage = _coverage(reconciliation)
    closed = sum(item.status == "closed_verified" for item in records.findings)
    lines = [
        f"# Executive report: {_md(records.engagement.name)}",
        "",
        "## Executive summary",
        "",
        f"This self-reviewed educational assessment is frozen to `{snapshot.id}`. "
        f"It contains workpapers for {coverage['workpapers_present']} of {coverage['total']} "
        f"controls, with executed conclusions for {coverage['assessed']}; limitation-only "
        "workpapers are not represented as tested.",
        "",
        "| Measure | Count |",
        "| --- | ---: |",
        f"| Registered controls | {coverage['total']} |",
        f"| Frozen workpapers | {coverage['workpapers_present']} |",
        f"| Executed assessment conclusions | {coverage['assessed']} |",
        f"| Findings | {len(records.findings)} |",
        f"| Actions | {len(records.actions)} |",
        f"| Retests | {len(records.retests)} |",
        f"| Verified closures | {closed} |",
        "",
        "## Finding summary",
        "",
        "| Finding | Controls | Rating | Status | Target date |",
        "| --- | --- | ---: | --- | --- |",
    ]
    lines.extend(
        f"| `{item.id}` | {_md(', '.join(sorted(item.control_ids)))} | {item.rating} | "
        f"`{item.status}` | {item.target_date} |"
        for item in sorted(records.findings, key=lambda value: value.id)
    )
    if not records.findings:
        lines.append("| None | - | - | - | - |")
    lines.extend(["", "## Limitations", ""])
    lines.extend(f"- {_md(item)}" for item in limitations)
    return "\n".join(lines) + "\n"


def _roadmap_rows(records: _FrozenRecords) -> list[dict[str, object]]:
    actions = {item.finding_id: item for item in records.actions}
    retests = {item.finding_id: item for item in records.retests}
    rows: list[dict[str, object]] = []
    for finding in sorted(records.findings, key=lambda item: item.id):
        action = actions.get(finding.id)
        retest = retests.get(finding.id)
        rows.append(
            {
                "finding_id": finding.id,
                "control_ids": ";".join(sorted(finding.control_ids)),
                "risk_rating": finding.rating,
                "finding_status": finding.status,
                "owner_role": finding.owner_role,
                "effort": finding.effort,
                "target_date": finding.target_date.isoformat(),
                "action_id": action.id if action else "",
                "action_status": action.status if action else "",
                "retest_id": retest.id if retest else "",
                "retest_result": retest.result if retest else "",
                "closure_decision": retest.decision if retest else "",
            }
        )
    return rows


def _csv_text(rows: Sequence[dict[str, object]]) -> str:
    fields = list(rows[0]) if rows else ["finding_id"]
    stream = io.StringIO(newline="")
    writer = csv.DictWriter(stream, fieldnames=fields, lineterminator="\n")
    writer.writeheader()
    writer.writerows(rows)
    return stream.getvalue()


def _append_table(
    sheet: Worksheet, headers: Sequence[str], rows: Iterable[Sequence[object]]
) -> None:
    sheet.append(list(headers))
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=_BLUE)
        cell.alignment = Alignment(vertical="top")
    for row in rows:
        sheet.append(list(row))
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for cells in sheet.columns:
        letter = cells[0].column_letter
        width = min(45, max(12, *(len(str(cell.value or "")) + 2 for cell in cells)))
        sheet.column_dimensions[letter].width = width
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)


def _cover(sheet: Worksheet, snapshot: ExportSnapshot, records: _FrozenRecords) -> None:
    sheet.sheet_view.showGridLines = False
    sheet.column_dimensions["A"].width = 28
    sheet.column_dimensions["B"].width = 88
    sheet["A1"] = "Cybersecurity Risk and ITGC Assessment"
    sheet["A1"].font = Font(size=18, bold=True, color=_BLUE)
    sheet.merge_cells("A1:B1")
    rows = (
        ("Engagement", records.engagement.name),
        ("Snapshot", snapshot.id),
        ("Frozen at", _iso(snapshot.created_at)),
        ("Review mode", records.engagement.review_mode),
        ("Boundary", "Synthetic local lab; self-reviewed; no assurance opinion"),
        ("PDF status", "Pending; HTML documents are print-ready"),
    )
    for index, (label, value) in enumerate(rows, start=3):
        sheet.cell(index, 1, label).font = Font(bold=True, color=_BLUE)
        sheet.cell(index, 2, value).alignment = Alignment(wrap_text=True, vertical="top")


def _normalize_xlsx(source: Path, destination: Path) -> None:
    """Rewrite OOXML members in stable order with fixed ZIP metadata."""
    with (
        zipfile.ZipFile(source, "r") as input_zip,
        zipfile.ZipFile(
            destination, "w", compression=zipfile.ZIP_DEFLATED, compresslevel=9
        ) as output_zip,
    ):
        for name in sorted(input_zip.namelist()):
            info = zipfile.ZipInfo(name, _ZIP_TIME)
            info.compress_type = zipfile.ZIP_DEFLATED
            info.external_attr = 0o600 << 16
            output_zip.writestr(info, input_zip.read(name))


def _build_workbook(
    path: Path,
    snapshot: ExportSnapshot,
    records: _FrozenRecords,
    reconciliation: dict[str, object],
    limitations: Sequence[str],
) -> None:
    workbook = Workbook()
    workbook.iso_dates = True
    workbook.properties.creator = "Cybersecurity Risk and ITGC Assessment"
    workbook.properties.title = records.engagement.name
    workbook.properties.created = snapshot.created_at.replace(tzinfo=None)
    workbook.properties.modified = snapshot.created_at.replace(tzinfo=None)
    cover = workbook.active
    cover.title = "Cover"
    _cover(cover, snapshot, records)

    scope = workbook.create_sheet("Scope")
    scope_rows = [("Scope", item) for item in records.engagement.scope]
    scope_rows.extend(("Exclusion", item) for item in records.engagement.exclusions)
    _append_table(scope, ("Type", "Statement"), scope_rows)

    assets = workbook.create_sheet("Assets")
    _append_table(
        assets,
        ("Asset ID", "Snapshot ID"),
        ((item, snapshot.id) for item in sorted(records.engagement.asset_ids)),
    )

    risk_controls: dict[str, list[str]] = {}
    for control in records.controls:
        for risk_id in control.risk_ids:
            risk_controls.setdefault(risk_id, []).append(control.id)
    risks = workbook.create_sheet("Risk Register")
    _append_table(
        risks,
        ("Risk ID", "Control IDs"),
        ((key, ";".join(sorted(value))) for key, value in sorted(risk_controls.items())),
    )

    workpaper_by_control = {item.control_id: item for item in records.workpapers}
    controls = workbook.create_sheet("Control Register")
    _append_table(
        controls,
        (
            "Control ID",
            "Domain",
            "Owner role",
            "Frequency",
            "Workpaper ID",
            "Design status",
            "Operating status",
        ),
        (
            (
                item.id,
                item.domain,
                item.owner_role,
                item.frequency,
                workpaper_by_control[item.id].id if item.id in workpaper_by_control else "",
                (
                    workpaper_by_control[item.id].design_status
                    if item.id in workpaper_by_control
                    else "not_assessed"
                ),
                (
                    workpaper_by_control[item.id].operating_status
                    if item.id in workpaper_by_control
                    else "not_tested"
                ),
            )
            for item in sorted(records.controls, key=lambda value: value.id)
        ),
    )

    workpapers = workbook.create_sheet("Workpapers")
    _append_table(
        workpapers,
        (
            "Workpaper ID",
            "Control ID",
            "Test type",
            "Evidence IDs",
            "Design status",
            "Operating status",
            "Review mode",
        ),
        (
            (
                item.id,
                item.control_id,
                item.test_type,
                ";".join(sorted(item.evidence_ids)),
                item.design_status,
                item.operating_status,
                item.review_mode,
            )
            for item in sorted(records.workpapers, key=lambda value: value.id)
        ),
    )

    evidence = workbook.create_sheet("Evidence Index")
    _append_table(
        evidence,
        ("Evidence ID", "Provenance", "Classification", "Media type", "Collected at"),
        (
            (item.id, item.provenance, item.sensitivity, item.media_type, _iso(item.collected_at))
            for item in sorted(records.evidence, key=lambda value: value.id)
            if item.sensitivity == "sanitized"
        ),
    )

    findings = workbook.create_sheet("Findings")
    _append_table(
        findings,
        (
            "Finding ID",
            "Control IDs",
            "Rating",
            "Status",
            "Evidence IDs",
            "Owner role",
            "Target date",
        ),
        (
            (
                item.id,
                ";".join(sorted(item.control_ids)),
                item.rating,
                item.status,
                ";".join(sorted(item.evidence_ids)),
                item.owner_role,
                item.target_date.isoformat(),
            )
            for item in sorted(records.findings, key=lambda value: value.id)
        ),
    )

    roadmap = _roadmap_rows(records)
    remediation = workbook.create_sheet("Remediation")
    _append_table(
        remediation,
        tuple(roadmap[0]) if roadmap else ("finding_id",),
        (tuple(row.values()) for row in roadmap),
    )

    retests = workbook.create_sheet("Retests")
    _append_table(
        retests,
        (
            "Retest ID",
            "Finding ID",
            "Action ID",
            "Evidence IDs",
            "Result",
            "Decision",
            "Executed at",
        ),
        (
            (
                item.id,
                item.finding_id,
                item.action_id,
                ";".join(sorted(item.evidence_ids)),
                item.result,
                item.decision,
                _iso(item.executed_at),
            )
            for item in sorted(records.retests, key=lambda value: value.id)
        ),
    )

    nist = workbook.create_sheet("NIST Mapping")
    _append_table(
        nist,
        ("Control ID", "NIST category"),
        (
            (item.id, category)
            for item in sorted(records.controls, key=lambda value: value.id)
            for category in sorted(item.nist_categories)
        ),
    )
    rmit = workbook.create_sheet("RMiT Mapping")
    _append_table(
        rmit,
        ("Control ID", "RMiT reference"),
        (
            (item.id, reference)
            for item in sorted(records.controls, key=lambda value: value.id)
            for reference in sorted(item.criteria_references)
            if reference.startswith("RMIT-")
        ),
    )

    counts = reconciliation["record_counts"]
    if not isinstance(counts, dict):
        raise TypeError("record counts must be a mapping")
    coverage = _coverage(reconciliation)
    reconcile = workbook.create_sheet("Reconciliation")
    reconciliation_rows = [(f"record_count.{key}", value) for key, value in sorted(counts.items())]
    reconciliation_rows.extend(
        (f"control_coverage.{key}", value) for key, value in sorted(coverage.items())
    )
    reconciliation_rows.extend([("snapshot_id", snapshot.id), ("unmatched_references", 0)])
    _append_table(reconcile, ("Measure", "Value"), reconciliation_rows)

    limits = workbook.create_sheet("Limitations")
    _append_table(
        limits,
        ("Number", "Limitation"),
        ((index, item) for index, item in enumerate(limitations, start=1)),
    )

    raw_path = path.with_suffix(".raw.xlsx")
    workbook.save(raw_path)
    _normalize_xlsx(raw_path, path)
    raw_path.unlink(missing_ok=True)


def _presentation_html(
    snapshot: ExportSnapshot,
    records: _FrozenRecords,
    reconciliation: dict[str, object],
    limitations: Sequence[str],
) -> str:
    coverage = _coverage(reconciliation)
    slides: list[tuple[str, list[str], str]] = [
        (
            "Assessment overview",
            [
                records.engagement.name,
                f"Snapshot {snapshot.id}",
                "Synthetic local lab · Self-reviewed",
            ],
            "Open by defining the educational scope and the evidence boundary.",
        ),
        (
            "Engagement scope",
            [
                *records.engagement.scope,
                (
                    f"Assessment period: {records.engagement.assessment_period.start} "
                    f"to {records.engagement.assessment_period.end}"
                ),
            ],
            "Explain what was assessed before results. State the exclusions explicitly.",
        ),
        (
            "Method",
            [
                "Frozen-snapshot source reconciliation",
                "Risk-to-control-to-evidence traceability",
                "Baseline, action, and retest chain",
                "Allowlisted sanitized reporting",
            ],
            "Describe the repeatable evidence chain. Do not imply external assurance.",
        ),
        (
            "Control coverage",
            [
                (
                    f"{coverage['workpapers_present']} of {coverage['total']} controls have "
                    "frozen workpapers"
                ),
                (
                    f"{coverage['assessed']} of {coverage['total']} controls have executed "
                    "assessment conclusions"
                ),
                f"Coverage state: {coverage['state']}",
                f"{coverage['unassessed']} controls remain unassessed",
            ],
            (
                "Lead with incomplete coverage so the register is not mistaken for a "
                "fully tested population."
            ),
        ),
        (
            "Findings and remediation",
            [
                f"Findings: {len(records.findings)}",
                f"Actions: {len(records.actions)}",
                f"Retests: {len(records.retests)}",
                (
                    "Verified closures: "
                    f"{sum(item.status == 'closed_verified' for item in records.findings)}"
                ),
            ],
            "Explain ratings and closure only at the level supported by frozen records.",
        ),
        (
            "Evidence and integrity",
            [
                (
                    "Sanitized evidence records: "
                    f"{sum(item.sensitivity == 'sanitized' for item in records.evidence)}"
                ),
                "Workbook contains evidence IDs, not evidence payloads",
                "Manifest hashes bind every published artifact",
                "No filesystem paths, queries, backups, or secrets are published",
            ],
            "Use evidence IDs for traceability without exposing raw lab material.",
        ),
        (
            "Limitations and next steps",
            [
                *limitations,
                (
                    "Next: complete remaining workpapers and obtain independent review "
                    "if assurance is required."
                ),
            ],
            "Close with limitations. PDF export is pending; the HTML deck is print-ready.",
        ),
    ]
    sections: list[str] = []
    for number, (title, bullets, notes) in enumerate(slides, start=1):
        items = "".join(f"<li>{html.escape(item)}</li>" for item in bullets)
        sections.append(
            f'<section class="slide" id="slide-{number}">'
            f'<div class="number">{number}/7</div><h1>{html.escape(title)}</h1>'
            f'<ul>{items}</ul><aside class="speaker-notes">'
            f"<strong>Speaker notes:</strong> {html.escape(notes)}</aside></section>"
        )
    return f"""<!doctype html>
<html lang="en"><head><meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>Assessment presentation - {html.escape(snapshot.id)}</title><style>
@page {{ size:16in 9in; margin:0 }} body {{ margin:0; background:#dfe7ef;
color:#172033; font-family:Arial,sans-serif }} .pdf-status {{ margin:16px auto;
max-width:1180px; background:#fff8c5; border:2px solid #9a6700; padding:10px }}
.slide {{ box-sizing:border-box; width:min(1200px,calc(100vw - 32px)); min-height:675px;
margin:24px auto; padding:64px 72px; background:white; border-top:12px solid #17365d;
position:relative; page-break-after:always }} .slide h1 {{ color:#17365d; font-size:40px }}
.slide li {{ font-size:24px; line-height:1.4; margin:14px 0 }} .number {{ position:absolute;
right:32px; top:24px; color:#566575 }} .speaker-notes {{ border-top:1px solid #aeb9c5;
color:#465566; font-size:14px; margin-top:36px; padding-top:12px }}
@media print {{ body {{ background:white }} .pdf-status,.speaker-notes {{ display:none }}
.slide {{ margin:0; width:100%; height:100vh }} }}
</style></head><body><p class="pdf-status"><strong>PDF status:</strong> Pending.
This seven-slide HTML deck is print-ready; no PDF is claimed or included.</p>
{"".join(sections)}</body></html>
"""


def _write_text(path: Path, value: str) -> None:
    path.write_text(value, encoding="utf-8", newline="\n")


def _manifest(snapshot: ExportSnapshot, directory: Path) -> dict[str, object]:
    artifacts = []
    for name in _ARTIFACT_NAMES:
        payload = (directory / name).read_bytes()
        artifacts.append(
            {
                "path": name,
                "sha256": hashlib.sha256(payload).hexdigest(),
                "size": len(payload),
                "classification": "sanitized_portfolio_artifact",
            }
        )
    return {
        "schema_version": "1.0",
        "snapshot_id": snapshot.id,
        "engagement_id": snapshot.engagement_id,
        "generated_at": _iso(snapshot.created_at),
        "artifacts": artifacts,
        "pdf": {
            "status": "pending",
            "reason": "No locked PDF renderer is available; print-ready HTML is provided.",
        },
    }


def create_final_artifacts(
    workspace: Path, snapshot_id: str, *, output_root: Path | None = None
) -> Path:
    """Render and atomically publish safe portfolio artifacts for one snapshot."""
    workspace = workspace.resolve()
    verify_snapshot(workspace, snapshot_id)
    snapshot = load_snapshot(workspace, snapshot_id)
    records = _load_frozen_records(workspace, snapshot)
    reconciliation = _reconcile(snapshot, records)
    limitations = _final_limitations(records, snapshot)
    export_root = (output_root or workspace / "outputs" / "artifacts").resolve()
    export_root.mkdir(parents=True, exist_ok=True)
    destination = export_root / snapshot.id
    if destination.exists():
        raise FileExistsError(f"artifact bundle already exists: {destination}")
    temporary = Path(tempfile.mkdtemp(prefix=f".{snapshot.id}.", dir=export_root))
    try:
        brief = _brief_markdown(snapshot, records, reconciliation, limitations)
        executive = _executive_markdown(snapshot, records, reconciliation, limitations)
        roadmap = _roadmap_rows(records)
        _write_text(temporary / "engagement-brief.md", brief)
        _write_text(
            temporary / "engagement-brief.html",
            _html_document("Engagement brief", brief),
        )
        _write_text(temporary / "executive-report.md", executive)
        _write_text(
            temporary / "executive-report.html",
            _html_document("Executive report", executive),
        )
        _write_text(temporary / "remediation-roadmap.csv", _csv_text(roadmap))
        _write_text(
            temporary / "remediation-roadmap.json",
            _json_text(
                {
                    "schema_version": "1.0",
                    "snapshot_id": snapshot.id,
                    "engagement_id": snapshot.engagement_id,
                    "items": roadmap,
                    "limitations": limitations,
                }
            ),
        )
        _write_text(
            temporary / "presentation.html",
            _presentation_html(snapshot, records, reconciliation, limitations),
        )
        _write_text(temporary / "reconciliation.json", _json_text(reconciliation))
        _write_text(
            temporary / "pdf-status.json",
            _json_text(
                {
                    "status": "pending",
                    "reason": (
                        "No locked PDF renderer is available; print-ready HTML is provided."
                    ),
                    "claimed_pdf_artifacts": [],
                }
            ),
        )
        _build_workbook(
            temporary / "assessment-workbook.xlsx",
            snapshot,
            records,
            reconciliation,
            limitations,
        )
        _write_text(
            temporary / "artifact-manifest.json",
            _json_text(_manifest(snapshot, temporary)),
        )
        verify_snapshot(workspace, snapshot_id)
        os.replace(temporary, destination)
    except BaseException:
        shutil.rmtree(temporary, ignore_errors=True)
        raise
    return destination
